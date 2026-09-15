# src/extraction/image_pipeline.py — v1
"""Embedded-image enrichment: extract bytes → dedup by sha256 → vision caption.

The per-format extractors (pdf/docx) only emit image *placeholders* (metadata);
the raw bytes are discarded. This module re-reads the document to pull the
actual image bytes, deduplicates them by content hash (so the same logo/figure
appearing N times is analysed once), captions each UNIQUE image via the vision
model, writes a per-image analysis artifact under `01_extraction/images/
img_{sha8}.json` (R-400-220 dedup layout), and folds the generated descriptions
back into `state.extraction_result.images` (replacing the placeholders).

The vision MODEL is resolved independently from the text agents: the caller
passes `llm_factory("image_analyzer")`, so a project can run a different
model/provider/size for image analysis than for summarisation.
"""

from __future__ import annotations

import hashlib
import io
import json
import logging
from dataclasses import dataclass
from typing import TYPE_CHECKING, Any

from ayextractor.core.models import ImageAnalysis
from ayextractor.extraction.image_analyzer import analyze_image
from ayextractor.storage.minio_layout import image_analysis_key

if TYPE_CHECKING:
    from ayextractor.pipeline.state import PipelineState
    from ayextractor.storage.base_output_writer import BaseOutputWriter
    from ayextractor.storage.minio_layout import RunPrefix

logger = logging.getLogger(__name__)


@dataclass(frozen=True)
class ImageBlob:
    """One embedded image with its raw bytes (the placeholder + the payload)."""

    image_id: str
    data: bytes
    media_type: str
    source_page: int | None


def sha256_hex(data: bytes) -> str:
    return hashlib.sha256(data).hexdigest()


def extract_image_blobs(raw_bytes: bytes, fmt: str) -> list[ImageBlob]:
    """Pull embedded image bytes from a document. PDF + DOCX only (the formats
    whose extractors emit image placeholders). Returns [] for other formats or
    when the optional dependency is missing (best-effort)."""
    fmt = (fmt or "").lower()
    try:
        if fmt == "pdf":
            return _pdf_blobs(raw_bytes)
        if fmt == "docx":
            return _docx_blobs(raw_bytes)
    except Exception:  # noqa: BLE001 — never let image extraction break a run
        logger.exception("Image byte extraction failed for format=%s", fmt)
    return []


def _pdf_blobs(raw_bytes: bytes) -> list[ImageBlob]:
    import fitz  # PyMuPDF

    blobs: list[ImageBlob] = []
    doc = fitz.open(stream=raw_bytes, filetype="pdf")
    try:
        for page_num in range(len(doc)):
            page = doc[page_num]
            for idx, info in enumerate(page.get_images(full=True)):
                xref = info[0]
                img_id = f"img_p{page_num + 1}_{idx + 1:02d}"
                try:
                    extracted = doc.extract_image(xref)
                except Exception:  # noqa: BLE001
                    logger.debug("extract_image failed for xref=%s", xref)
                    continue
                data = extracted.get("image")
                if not data:
                    continue
                media = f"image/{extracted.get('ext', 'png')}"
                blobs.append(ImageBlob(img_id, data, media, page_num + 1))
    finally:
        doc.close()
    return blobs


def _docx_blobs(raw_bytes: bytes) -> list[ImageBlob]:
    import docx

    blobs: list[ImageBlob] = []
    document = docx.Document(io.BytesIO(raw_bytes))
    count = 0
    for rel in document.part.rels.values():
        if "image" not in rel.reltype:
            continue
        count += 1
        try:
            part = rel.target_part
            data = part.blob
            media = getattr(part, "content_type", None) or "image/png"
        except Exception:  # noqa: BLE001
            logger.debug("docx image blob read failed")
            continue
        if data:
            blobs.append(ImageBlob(f"img_{count:03d}", data, media, None))
    return blobs


# Images smaller than this (raw encoded bytes) are almost always decorative —
# icons, logos, bullets, rules, separators. Captioning them wastes a vision
# call for no retrieval value, so they are skipped. Cheap heuristic, no image
# decoding needed (so no Pillow dependency).
_MIN_CAPTION_BYTES = 3072


async def caption_image_blobs(
    *,
    blobs: list[ImageBlob],
    state: PipelineState,
    llm_factory: Any,
    writer: BaseOutputWriter | None = None,
    run_prefix: RunPrefix | None = None,
) -> dict[str, ImageAnalysis]:
    """Deduplicate `blobs` by sha256, caption each UNIQUE image via the vision
    model, write a per-image JSON artifact, and fold the descriptions back into
    `state.extraction_result.images`. Returns {sha256: ImageAnalysis}.

    Best-effort: a single image's failure is recorded in `state.errors` and the
    rest continue. Identical images are captioned ONCE (dedup)."""
    if not blobs:
        return {}

    by_sha: dict[str, list[ImageBlob]] = {}
    for blob in blobs:
        by_sha.setdefault(sha256_hex(blob.data), []).append(blob)

    llm = llm_factory("image_analyzer")
    analyses: dict[str, ImageAnalysis] = {}
    for sha, group in by_sha.items():
        rep = group[0]
        if len(rep.data) < _MIN_CAPTION_BYTES:
            logger.debug(
                "skipping decorative image %s (%d bytes < %d)",
                sha[:8], len(rep.data), _MIN_CAPTION_BYTES,
            )
            continue
        try:
            analysis = await analyze_image(
                rep.data, rep.media_type, f"img_{sha[:8]}", llm, rep.source_page
            )
        except Exception as exc:  # noqa: BLE001
            logger.warning("image analysis failed for %s: %s", sha[:8], exc)
            state.errors.append(f"image_analyzer:{sha[:8]}: {exc}")
            continue
        analyses[sha] = analysis
        state.total_llm_calls += 1
        if writer is not None and run_prefix is not None:
            payload = {
                "sha256": sha,
                **analysis.model_dump(),
                "occurrences": len(group),
                "image_ids": [b.image_id for b in group],
                "source_pages": [b.source_page for b in group],
            }
            await writer.write(
                image_analysis_key(run_prefix, sha[:8]), json.dumps(payload, indent=2)
            )

    _fold_into_extraction(state, by_sha, analyses)
    return analyses


def _fold_into_extraction(
    state: PipelineState,
    by_sha: dict[str, list[ImageBlob]],
    analyses: dict[str, ImageAnalysis],
) -> None:
    """Replace placeholder descriptions in `state.extraction_result.images` with
    the generated captions (matched by the shared `image_id` scheme)."""
    extraction = getattr(state, "extraction_result", None)
    if extraction is None or not extraction.images:
        return
    sha_for_id: dict[str, str] = {}
    for sha, group in by_sha.items():
        for blob in group:
            sha_for_id[blob.image_id] = sha
    folded: list[ImageAnalysis] = []
    for img in extraction.images:
        sha = sha_for_id.get(img.id)
        analysis = analyses.get(sha) if sha else None
        if analysis is not None:
            folded.append(
                ImageAnalysis(
                    id=img.id,
                    type=analysis.type,
                    description=analysis.description,
                    entities=analysis.entities,
                    source_page=img.source_page,
                )
            )
        else:
            folded.append(img)
    extraction.images = folded


async def run_image_enrichment(
    *,
    raw_bytes: bytes,
    fmt: str,
    state: PipelineState,
    llm_factory: Any,
    writer: BaseOutputWriter | None = None,
    run_prefix: RunPrefix | None = None,
) -> None:
    """Extract image bytes from the document then dedup + caption them."""
    blobs = extract_image_blobs(raw_bytes, fmt)
    if not blobs:
        return
    await caption_image_blobs(
        blobs=blobs,
        state=state,
        llm_factory=llm_factory,
        writer=writer,
        run_prefix=run_prefix,
    )
