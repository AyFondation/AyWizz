# tests/e2e/test_func_extraction.py — v2
"""Functional E2E tests — document extraction across formats.

Changelog:
    v2: Add conditional skip for DOCX tests when python-docx is not installed.
    v1: Initial version.

Validates that extractors produce correct output for known input documents.
Uses golden files (tests/e2e/fixtures/) and expected output schemas
(tests/e2e/expected/) for deterministic, repeatable assertions.

No LLM required — these test pure extraction logic.

Coverage:
    - PDF extraction (text, table, images, structure)
    - DOCX extraction (text, table, images, structure)
    - Cross-format consistency (PDF vs DOCX same content)
    - Image-as-document input (PNG → placeholder)
    - Error handling (corrupted/empty input)
"""

from __future__ import annotations

from pathlib import Path

import pytest

pytestmark = [pytest.mark.e2e]


def _skip_if_no_docx():
    """Skip test if python-docx is not installed."""
    try:
        import docx  # noqa: F401
    except ImportError:
        pytest.skip("python-docx not installed — pip install python-docx")


# =====================================================================
#  PDF EXTRACTION
# =====================================================================

class TestPdfExtraction:
    """PDF extractor against golden NIS2 test document."""

    @pytest.mark.asyncio
    async def test_extract_returns_result(self, nis2_pdf_path, expected_extraction):
        """PDF extraction returns a valid ExtractionResult."""
        from ayextractor.extraction.pdf_extractor import PdfExtractor

        extractor = PdfExtractor()
        result = await extractor.extract(nis2_pdf_path)

        assert result is not None
        assert result.raw_text is not None
        assert len(result.raw_text) >= expected_extraction["min_text_length"]

    @pytest.mark.asyncio
    async def test_extract_from_bytes(self, nis2_pdf_bytes, expected_extraction):
        """PDF extraction works from raw bytes (not just file path)."""
        from ayextractor.extraction.pdf_extractor import PdfExtractor

        extractor = PdfExtractor()
        result = await extractor.extract(nis2_pdf_bytes)

        assert len(result.raw_text) >= expected_extraction["min_text_length"]

    @pytest.mark.asyncio
    async def test_text_contains_key_content(self, nis2_pdf_path, expected_extraction):
        """Extracted text contains all expected key fragments."""
        from ayextractor.extraction.pdf_extractor import PdfExtractor

        extractor = PdfExtractor()
        result = await extractor.extract(nis2_pdf_path)

        text = result.raw_text
        for fragment in expected_extraction["must_contain_text"]:
            assert fragment in text, (
                f"Expected text fragment not found: {fragment!r}"
            )

    @pytest.mark.asyncio
    async def test_structure_detection(self, nis2_pdf_path, expected_extraction):
        """Extracted structure contains expected sections."""
        from ayextractor.extraction.pdf_extractor import PdfExtractor

        extractor = PdfExtractor()
        result = await extractor.extract(nis2_pdf_path)

        assert result.structure is not None
        section_titles = [
            s.title.lower() for s in (result.structure.sections or [])
        ]

        for expected_section in expected_extraction["must_contain_sections"]:
            title_fragment = expected_section["title_contains"].lower()
            found = any(title_fragment in t for t in section_titles)
            assert found, (
                f"Expected section containing {title_fragment!r} not found. "
                f"Got sections: {section_titles}"
            )

    @pytest.mark.asyncio
    async def test_table_extraction(self, nis2_pdf_path, expected_extraction):
        """PDF table on page 2 is extracted with correct content."""
        from ayextractor.extraction.pdf_extractor import PdfExtractor

        extractor = PdfExtractor()
        result = await extractor.extract(nis2_pdf_path)

        assert len(result.tables) >= expected_extraction["min_tables"], (
            f"Expected at least {expected_extraction['min_tables']} table(s), "
            f"got {len(result.tables)}"
        )

        # Flatten all table content for assertion
        all_table_text = " ".join(t.content_markdown for t in result.tables)
        for fragment in expected_extraction["table_must_contain"]:
            assert fragment in all_table_text, (
                f"Table content missing: {fragment!r}"
            )

    @pytest.mark.asyncio
    async def test_image_metadata(self, nis2_pdf_path):
        """PDF extraction detects embedded images (metadata only)."""
        from ayextractor.extraction.pdf_extractor import PdfExtractor

        extractor = PdfExtractor()
        result = await extractor.extract(nis2_pdf_path)

        # Our test PDF has no raster images embedded (diagram is vector),
        # so images list depends on PyMuPDF detection.
        # We just verify the field is a list.
        assert isinstance(result.images, list)

    @pytest.mark.asyncio
    async def test_language_detection(self, nis2_pdf_path, expected_extraction):
        """Extraction reports correct language."""
        from ayextractor.extraction.pdf_extractor import PdfExtractor

        extractor = PdfExtractor()
        result = await extractor.extract(nis2_pdf_path)

        assert result.language == expected_extraction["expected_language"]


# =====================================================================
#  DOCX EXTRACTION
# =====================================================================

class TestDocxExtraction:
    """DOCX extractor against golden NIS2 test document."""

    @pytest.mark.asyncio
    async def test_extract_returns_result(self, nis2_docx_path, expected_extraction):
        """DOCX extraction returns a valid ExtractionResult."""
        _skip_if_no_docx()
        from ayextractor.extraction.docx_extractor import DocxExtractor

        extractor = DocxExtractor()
        result = await extractor.extract(nis2_docx_path)

        assert result is not None
        assert len(result.raw_text) >= expected_extraction["min_text_length"]

    @pytest.mark.asyncio
    async def test_extract_from_bytes(self, nis2_docx_bytes, expected_extraction):
        """DOCX extraction works from raw bytes."""
        _skip_if_no_docx()
        from ayextractor.extraction.docx_extractor import DocxExtractor

        extractor = DocxExtractor()
        result = await extractor.extract(nis2_docx_bytes)

        assert len(result.raw_text) >= expected_extraction["min_text_length"]

    @pytest.mark.asyncio
    async def test_text_contains_key_content(self, nis2_docx_path, expected_extraction):
        """Extracted text contains all expected key fragments."""
        _skip_if_no_docx()
        from ayextractor.extraction.docx_extractor import DocxExtractor

        extractor = DocxExtractor()
        result = await extractor.extract(nis2_docx_path)

        text = result.raw_text
        for fragment in expected_extraction["must_contain_text"]:
            assert fragment in text, (
                f"Expected text fragment not found in DOCX: {fragment!r}"
            )

    @pytest.mark.asyncio
    async def test_structure_detection(self, nis2_docx_path, expected_extraction):
        """DOCX structure detection finds expected sections."""
        _skip_if_no_docx()
        from ayextractor.extraction.docx_extractor import DocxExtractor

        extractor = DocxExtractor()
        result = await extractor.extract(nis2_docx_path)

        assert result.structure is not None
        section_titles = [
            s.title.lower() for s in (result.structure.sections or [])
        ]

        for expected_section in expected_extraction["must_contain_sections"]:
            title_fragment = expected_section["title_contains"].lower()
            found = any(title_fragment in t for t in section_titles)
            assert found, (
                f"Expected section containing {title_fragment!r} not found in DOCX. "
                f"Got sections: {section_titles}"
            )

    @pytest.mark.asyncio
    async def test_table_extraction(self, nis2_docx_path, expected_extraction):
        """DOCX table extraction captures compliance data."""
        _skip_if_no_docx()
        from ayextractor.extraction.docx_extractor import DocxExtractor

        extractor = DocxExtractor()
        result = await extractor.extract(nis2_docx_path)

        assert len(result.tables) >= expected_extraction["min_tables"]

        all_table_text = " ".join(t.content_markdown for t in result.tables)
        for fragment in expected_extraction["table_must_contain"]:
            assert fragment in all_table_text, (
                f"DOCX table content missing: {fragment!r}"
            )


# =====================================================================
#  CROSS-FORMAT CONSISTENCY
# =====================================================================

class TestCrossFormatConsistency:
    """Same document in PDF vs DOCX must produce semantically equivalent output."""

    @pytest.mark.asyncio
    async def test_key_entities_present_in_both(
        self, nis2_pdf_path, nis2_docx_path, expected_extraction
    ):
        """Both formats must contain the same key text fragments."""
        _skip_if_no_docx()
        from ayextractor.extraction.pdf_extractor import PdfExtractor
        from ayextractor.extraction.docx_extractor import DocxExtractor

        pdf_result = await PdfExtractor().extract(nis2_pdf_path)
        docx_result = await DocxExtractor().extract(nis2_docx_path)

        for fragment in expected_extraction["must_contain_text"]:
            assert fragment in pdf_result.raw_text, (
                f"PDF missing: {fragment!r}"
            )
            assert fragment in docx_result.raw_text, (
                f"DOCX missing: {fragment!r}"
            )

    @pytest.mark.asyncio
    async def test_table_data_consistent(
        self, nis2_pdf_path, nis2_docx_path, expected_extraction
    ):
        """Both formats extract the same table content."""
        _skip_if_no_docx()
        from ayextractor.extraction.pdf_extractor import PdfExtractor
        from ayextractor.extraction.docx_extractor import DocxExtractor

        pdf_result = await PdfExtractor().extract(nis2_pdf_path)
        docx_result = await DocxExtractor().extract(nis2_docx_path)

        # Both should have at least one table
        assert len(pdf_result.tables) >= 1
        assert len(docx_result.tables) >= 1

        # Both tables should contain the same key values
        for fragment in expected_extraction["table_must_contain"]:
            pdf_tables_text = " ".join(t.content_markdown for t in pdf_result.tables)
            docx_tables_text = " ".join(t.content_markdown for t in docx_result.tables)
            assert fragment in pdf_tables_text, f"PDF table missing: {fragment!r}"
            assert fragment in docx_tables_text, f"DOCX table missing: {fragment!r}"

    @pytest.mark.asyncio
    async def test_section_count_comparable(self, nis2_pdf_path, nis2_docx_path):
        """Both formats detect a similar number of sections."""
        _skip_if_no_docx()
        from ayextractor.extraction.pdf_extractor import PdfExtractor
        from ayextractor.extraction.docx_extractor import DocxExtractor

        pdf_result = await PdfExtractor().extract(nis2_pdf_path)
        docx_result = await DocxExtractor().extract(nis2_docx_path)

        pdf_sections = len(pdf_result.structure.sections) if pdf_result.structure else 0
        docx_sections = len(docx_result.structure.sections) if docx_result.structure else 0

        # Allow some variance but both should detect sections
        assert pdf_sections >= 2, f"PDF detected only {pdf_sections} sections"
        assert docx_sections >= 2, f"DOCX detected only {docx_sections} sections"


# =====================================================================
#  IMAGE-AS-DOCUMENT INPUT
# =====================================================================

class TestImageInputExtraction:
    """Image file submitted as a document (not embedded image)."""

    @pytest.mark.asyncio
    async def test_png_returns_extraction_result(
        self, nis2_table_png_path, expected_image_input
    ):
        """PNG input returns a valid ExtractionResult with image metadata."""
        from ayextractor.extraction.image_input_extractor import ImageInputExtractor

        extractor = ImageInputExtractor()
        result = await extractor.extract(nis2_table_png_path)

        assert result is not None
        assert expected_image_input["must_return_extraction_result"]

        # Should have placeholder text (pending LLM Vision)
        assert len(result.raw_text) > 0

        # Should register the image for later analysis
        assert len(result.images) >= expected_image_input["min_images"]

    @pytest.mark.asyncio
    async def test_diagram_png_extraction(self, nis2_diagram_png_path):
        """Diagram PNG is accepted and returns valid result."""
        from ayextractor.extraction.image_input_extractor import ImageInputExtractor

        extractor = ImageInputExtractor()
        result = await extractor.extract(nis2_diagram_png_path)

        assert result is not None
        assert len(result.images) >= 1
        assert "image" in result.images[0].description.lower() or result.images[0].id

    @pytest.mark.asyncio
    async def test_image_from_bytes(self, nis2_table_png_path):
        """Image extraction works from raw bytes."""
        from ayextractor.extraction.image_input_extractor import ImageInputExtractor

        extractor = ImageInputExtractor()
        img_bytes = nis2_table_png_path.read_bytes()
        result = await extractor.extract(img_bytes)

        assert result is not None
        assert len(result.raw_text) > 0


# =====================================================================
#  ERROR HANDLING
# =====================================================================

class TestExtractionErrors:
    """Extractors handle bad input gracefully."""

    @pytest.mark.asyncio
    async def test_pdf_nonexistent_file(self):
        """PDF extractor raises on non-existent file."""
        from ayextractor.extraction.pdf_extractor import PdfExtractor

        extractor = PdfExtractor()
        with pytest.raises((FileNotFoundError, Exception)):
            await extractor.extract(Path("/nonexistent/fake.pdf"))

    @pytest.mark.asyncio
    async def test_docx_nonexistent_file(self):
        """DOCX extractor raises on non-existent file."""
        _skip_if_no_docx()
        from ayextractor.extraction.docx_extractor import DocxExtractor

        extractor = DocxExtractor()
        with pytest.raises((FileNotFoundError, Exception)):
            await extractor.extract(Path("/nonexistent/fake.docx"))

    @pytest.mark.asyncio
    async def test_image_nonexistent_file(self):
        """Image extractor raises on non-existent file."""
        from ayextractor.extraction.image_input_extractor import ImageInputExtractor

        extractor = ImageInputExtractor()
        with pytest.raises(FileNotFoundError):
            await extractor.extract(Path("/nonexistent/fake.png"))

    @pytest.mark.asyncio
    async def test_pdf_corrupted_bytes(self):
        """PDF extractor raises on corrupted bytes."""
        from ayextractor.extraction.pdf_extractor import PdfExtractor

        extractor = PdfExtractor()
        with pytest.raises(Exception):
            await extractor.extract(b"this is not a valid PDF")

    @pytest.mark.asyncio
    async def test_docx_corrupted_bytes(self):
        """DOCX extractor raises on corrupted bytes."""
        _skip_if_no_docx()
        from ayextractor.extraction.docx_extractor import DocxExtractor

        extractor = DocxExtractor()
        with pytest.raises(Exception):
            await extractor.extract(b"this is not a valid DOCX")